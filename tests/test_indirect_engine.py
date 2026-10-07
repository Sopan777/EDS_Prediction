"""Tests for the standalone Indirect Material Source Rule Engine (`indirect_engine`)."""
from __future__ import annotations

import json
import os
import openpyxl
import pytest
from django.test import Client

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
import django
django.setup()

from indirect_engine.reference_loader import (
    EXCEL_PATH,
    compute_tolerance_band,
    load_indirect_reference,
)
from indirect_engine.engine import (
    classify_indirect_material_family,
    predict_indirect_particle,
    predict_indirect_source,
)


def test_tolerance_rules_exact_formula():
    """Verify exact tolerance rules: <1 -> ±25%, 1..5 -> ±20%, >5 -> ±10%."""
    # < 1 -> ±25%
    tol_sub1, min_sub1, max_sub1 = compute_tolerance_band(0.8)
    assert tol_sub1 == 0.25
    assert min_sub1 == pytest.approx(0.8 * 0.75, abs=1e-6)
    assert max_sub1 == pytest.approx(0.8 * 1.25, abs=1e-6)

    # Boundary = 1.0 -> 1 to 5 inclusive -> ±20%
    tol_one, min_one, max_one = compute_tolerance_band(1.0)
    assert tol_one == 0.20
    assert min_one == pytest.approx(0.80, abs=1e-6)
    assert max_one == pytest.approx(1.20, abs=1e-6)

    # Inside 1..5 -> ±20%
    tol_mid, min_mid, max_mid = compute_tolerance_band(1.56)
    assert tol_mid == 0.20
    assert min_mid == pytest.approx(1.56 * 0.80, abs=1e-6)
    assert max_mid == pytest.approx(1.56 * 1.20, abs=1e-6)

    # Boundary = 5.0 -> 1 to 5 inclusive -> ±20%
    tol_five, min_five, max_five = compute_tolerance_band(5.0)
    assert tol_five == 0.20
    assert min_five == pytest.approx(4.0, abs=1e-6)
    assert max_five == pytest.approx(6.0, abs=1e-6)

    # > 5 -> ±10%
    tol_gt5, min_gt5, max_gt5 = compute_tolerance_band(18.21)
    assert tol_gt5 == 0.10
    assert min_gt5 == pytest.approx(18.21 * 0.90, abs=1e-6)
    assert max_gt5 == pytest.approx(18.21 * 1.10, abs=1e-6)


def test_cleaning_area_excel_min_max_sheet_agreement():
    """Verify our tolerance computation matches all 167 populated elemental values in Cleaning area.xlsx."""
    assert EXCEL_PATH.exists(), "Cleaning area.xlsx must exist in data/"
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Min Max Values"]

    checked = 0
    for r in range(3, ws.max_row + 1):
        sn = ws.cell(r, 1).value
        ref_val = ws.cell(r, 5).value
        tol_val = ws.cell(r, 6).value
        min_v = ws.cell(r, 7).value
        max_v = ws.cell(r, 8).value
        if sn is None or ref_val is None:
            continue
        tol, calc_min, calc_max = compute_tolerance_band(float(ref_val))
        assert tol == pytest.approx(float(tol_val), abs=1e-4)
        assert calc_min == pytest.approx(float(min_v), abs=1e-4)
        assert calc_max == pytest.approx(float(max_v), abs=1e-4)
        checked += 1

    assert checked > 100


def test_blanks_not_treated_as_zero():
    """Ensure blank elemental cells in Cleaning area.xlsx are omitted, never stored as 0.0."""
    parts = load_indirect_reference(force_reload=True)
    assert len(parts) == 24
    for part in parts:
        assert len(part.elements) > 0
        for sym, band in part.elements.items():
            assert band.reference_value is not None
            assert band.reference_value > 0.0, f"{part.part_name} has 0.0 for {sym}"


def test_all_24_cleaning_area_parts_two_stage_self_consistency():
    """Every reference part in Cleaning area.xlsx must predict its own Material Family (Stage 1) and rank #1 at 100% (Stage 2)."""
    parts = load_indirect_reference()
    for part in parts:
        raw_comp = {el: band.reference_value for el, band in part.elements.items()}
        if "Fe" not in raw_comp:
            rem = 100.0 - sum(
                v for k, v in raw_comp.items() if k not in ("C", "O", "N")
            )
            raw_comp["Fe"] = max(0.0, round(rem, 2))

        pred = predict_indirect_source(raw_comp)
        assert pred["indirectFamily"] == part.material, (
            f"Stage 1 mismatch for {part.part_name}: got {pred['indirectFamily']}, expected {part.material}"
        )
        top = pred["topIndirectSource"]
        assert top is not None
        assert top["partName"] == part.part_name, (
            f"Stage 2 mismatch for {part.part_name}: got {top['partName']}"
        )
        assert top["compatibilityPct"] == 100
        assert len(top["outsideToleranceElements"]) == 0


def test_stage_separation_and_unknown_inconclusive():
    """Verify non-matching / contradictory compositions return Unknown / Inconclusive instead of forcing a match."""
    # Cu-Sn Bronze (non-ferrous) -> no Cleaning Area indirect family matches
    bronze_pred = predict_indirect_source({"Cu": 89.2, "Sn": 7.6, "P": 0.3})
    assert bronze_pred["decision"] == "unknown"
    assert bronze_pred["indirectFamily"] == "Unknown / Inconclusive"
    assert bronze_pred["statusLabel"] == "Unknown / Inconclusive"
    assert bronze_pred["topIndirectSource"] is None

    # High-Mo / High-Cr alloy that contradicts all 6 Cleaning Area families
    contradictory = predict_indirect_source({"Fe": 65.0, "Cr": 5.5, "Mo": 12.0, "W": 8.0})
    assert contradictory["decision"] == "unknown"
    assert contradictory["statusLabel"] == "Unknown / Inconclusive"
    assert contradictory["topIndirectSource"] is None


def test_indirect_predict_api_and_analyze_integration():
    """Verify /api/indirect-predict and /api/analyze return independent indirectSourcePrediction payloads."""
    client = Client()

    # 1. Direct call to /api/indirect-predict with IC Stud Tray (SS 304) composition
    ic_stud_tray_spec = {
        "Mn": 1.51,
        "Cu": 0.34,
        "Ni": 8.08,
        "Cr": 18.17,
        "V": 0.08,
        "Co": 0.063,
        "Mo": 0.14,
        "Fe": 71.617,
    }
    resp = client.post(
        "/api/indirect-predict",
        data=json.dumps({"spectra": [ic_stud_tray_spec]}),
        content_type="application/json",
    )
    assert resp.status_code == 200
    ind_data = resp.json()
    assert ind_data["indirectFamily"] == "SS 304"
    assert ind_data["topIndirectSource"]["partName"] == "IC Stud Tray"
    assert ind_data["topIndirectSource"]["compatibilityPct"] == 100
    assert ind_data["decision"] in ("identified", "ambiguous")

    # 2. Call /api/analyze and verify both direct and indirect predictions exist independently
    valve_set_base_spec = {
        "Mn": 0.69,
        "Cu": 0.25,
        "Ni": 0.14,
        "Cr": 1.41,
        "Mo": 0.056,
        "Fe": 97.454,
    }
    resp_analyze = client.post(
        "/api/analyze",
        data=json.dumps({
            "spectra": [
                ic_stud_tray_spec,
                valve_set_base_spec,
            ]
        }),
        content_type="application/json",
    )
    assert resp_analyze.status_code == 200
    full_data = resp_analyze.json()

    # Top-level pooled indirect prediction exists
    assert "indirectSourcePrediction" in full_data
    assert len(full_data["perSpectrum"]) == 2

    # Spectrum 1: Direct engine -> F4 (Austenitic SS), Indirect engine -> SS 304 -> IC Stud Tray
    spec1 = full_data["perSpectrum"][0]
    assert spec1["familyCode"] == "F4"
    assert spec1["indirectSourcePrediction"]["indirectFamily"] == "SS 304"
    assert spec1["indirectSourcePrediction"]["topIndirectSource"]["partName"] == "IC Stud Tray"

    # Spectrum 2: Direct engine runs its own rules independently while Indirect engine predicts EN 31 -> Valve Set removing Base
    spec2 = full_data["perSpectrum"][1]
    assert spec2["familyCode"] in ("F1a", "F2")
    assert spec2["indirectSourcePrediction"]["indirectFamily"] == "EN 31"
    assert spec2["indirectSourcePrediction"]["topIndirectSource"]["partName"] == "Valve Set removing Base"

    # 3. Verify /analyzer/results/ and /analyzer/ templates include the Probable Indirect Source section
    res_page = client.get("/analyzer/results/")
    assert res_page.status_code == 200
    html = res_page.content.decode("utf-8")
    assert "Probable Indirect Source" in html

    res_index = client.get("/analyzer/")
    assert res_index.status_code == 200
    assert "Probable Indirect Source" in res_index.content.decode("utf-8")

    # 4. Verify the full multi-spectrum result was saved to AnalysisHistory in the database
    from apps.history.models import AnalysisHistory

    analysis_id = full_data.get("analysisId")
    assert analysis_id is not None
    saved_rec = AnalysisHistory.objects.filter(id=analysis_id).first()
    assert saved_rec is not None
    saved_full = saved_rec.get_full_result()
    assert saved_full is not None
    assert len(saved_full["perSpectrum"]) == 2
    assert saved_full["perSpectrum"][0]["indirectSourcePrediction"]["topIndirectSource"]["partName"] == "IC Stud Tray"
    assert saved_full["perSpectrum"][1]["indirectSourcePrediction"]["topIndirectSource"]["partName"] == "Valve Set removing Base"

    # 5. Verify /history/?id=<analysis_id> renders the exact Results UI with the saved prediction
    res_hist = client.get(f"/history/?id={analysis_id}")
    assert res_hist.status_code == 200
    hist_html = res_hist.content.decode("utf-8")
    assert "All Extracted Spectra — Prediction Overview" in hist_html
    assert "Probable Indirect Source" in hist_html
    assert "Detailed Per-Spectrum Predictions" in hist_html
    assert analysis_id in hist_html

