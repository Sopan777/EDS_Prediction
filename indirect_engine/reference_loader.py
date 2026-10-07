"""
indirect_engine/reference_loader.py
===================================
Loads the Indirect Material Source reference dataset from `data/Cleaning area.xlsx`
(or cached `indirect_engine/knowledge/cleaning_area_reference.json`) and computes
the exact element tolerance ranges:
  - Reference Value < 1           -> Tolerance ±25% (Min = V * 0.75, Max = V * 1.25)
  - 1 <= Reference Value <= 5     -> Tolerance ±20% (Min = V * 0.80, Max = V * 1.20)
  - Reference Value > 5           -> Tolerance ±10% (Min = V * 0.90, Max = V * 1.10)

Blank / missing elemental reference values are NEVER converted to zero.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


ROOT_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT_DIR / "data" / "Cleaning area.xlsx"
JSON_CACHE_PATH = Path(__file__).resolve().parent / "knowledge" / "cleaning_area_reference.json"


def compute_tolerance_band(ref_value: float) -> Tuple[float, float, float]:
    """
    Calculate (tolerance_fraction, min_value, max_value) for a populated reference value:
      - If reference value < 1           -> tolerance ±25% (0.25)
      - If reference value 1 to 5 (incl) -> tolerance ±20% (0.20)
      - If reference value > 5           -> tolerance ±10% (0.10)
    """
    val = float(ref_value)
    if val < 1.0:
        tol = 0.25
    elif val <= 5.0:
        tol = 0.20
    else:
        tol = 0.10
    min_val = round(val * (1.0 - tol), 6)
    max_val = round(val * (1.0 + tol), 6)
    return tol, min_val, max_val


@dataclass(frozen=True)
class IndirectElementBand:
    element: str
    reference_value: float
    tolerance: float
    min_value: float
    max_value: float

    @property
    def tolerance_pct_label(self) -> str:
        return f"±{int(round(self.tolerance * 100))}%"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "element": self.element,
            "referenceValue": self.reference_value,
            "tolerance": self.tolerance,
            "toleranceLabel": self.tolerance_pct_label,
            "minValue": round(self.min_value, 4),
            "maxValue": round(self.max_value, 4),
        }


@dataclass(frozen=True)
class IndirectPartReference:
    sn: int
    part_name: str
    location: str
    material: str
    elements: Dict[str, IndirectElementBand] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "sn": self.sn,
            "partName": self.part_name,
            "location": self.location,
            "material": self.material,
            "elements": {k: v.to_dict() for k, v in self.elements.items()},
        }


_CACHED_PARTS: Optional[List[IndirectPartReference]] = None


def _clean_element_header(raw_header: Any) -> Optional[str]:
    if raw_header is None:
        return None
    s = str(raw_header).strip()
    if not s:
        return None
    s = s.replace("(%)", "").replace("%", "").strip()
    return s or None


def build_reference_from_excel(excel_path: Path = EXCEL_PATH) -> Dict[str, Any]:
    """Parse `Cleaning area.xlsx` and build the canonical indirect reference dictionary."""
    import openpyxl

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Cleaning Area"]

    headers: List[Optional[str]] = [
        _clean_element_header(ws.cell(2, col).value) if col >= 6 else str(ws.cell(2, col).value or "").strip()
        for col in range(1, ws.max_column + 1)
    ]

    parts_data: List[Dict[str, Any]] = []
    families_summary: Dict[str, List[str]] = {}

    for row_idx in range(3, ws.max_row + 1):
        sn_raw = ws.cell(row_idx, 1).value
        part_raw = ws.cell(row_idx, 2).value
        loc_raw = ws.cell(row_idx, 4).value
        mat_raw = ws.cell(row_idx, 5).value

        if sn_raw is None or part_raw is None or mat_raw is None:
            continue

        sn = int(sn_raw)
        part_name = str(part_raw).strip()
        location = str(loc_raw or "").strip()
        material = str(mat_raw).strip()

        elements_dict: Dict[str, Dict[str, Any]] = {}
        for col_idx in range(6, ws.max_column + 1):
            el_sym = headers[col_idx - 1]
            if not el_sym:
                continue
            cell_val = ws.cell(row_idx, col_idx).value
            # Strictly skip blank/missing cells — never treat blank as zero
            if cell_val is None or str(cell_val).strip() == "":
                continue
            try:
                ref_val = float(cell_val)
            except (ValueError, TypeError):
                continue

            tol, min_v, max_v = compute_tolerance_band(ref_val)
            elements_dict[el_sym] = {
                "element": el_sym,
                "reference_value": ref_val,
                "tolerance": tol,
                "min_value": min_v,
                "max_value": max_v,
            }

        parts_data.append({
            "sn": sn,
            "part_name": part_name,
            "location": location,
            "material": material,
            "elements": elements_dict,
        })
        families_summary.setdefault(material, []).append(part_name)

    payload = {
        "version": "1.0.0",
        "source_file": "data/Cleaning area.xlsx",
        "scope": "Indirect Material Composition - Cleaning Area",
        "tolerance_rules": [
            {"range": "< 1", "tolerance": 0.25, "min_formula": "Value * 0.75", "max_formula": "Value * 1.25"},
            {"range": "1 to 5 (inclusive)", "tolerance": 0.20, "min_formula": "Value * 0.80", "max_formula": "Value * 1.20"},
            {"range": "> 5", "tolerance": 0.10, "min_formula": "Value * 0.90", "max_formula": "Value * 1.10"},
        ],
        "material_families": families_summary,
        "parts": parts_data,
    }

    JSON_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(JSON_CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2)

    return payload


def load_indirect_reference(
    force_reload: bool = False,
    rebuild_from_excel: bool = False,
) -> List[IndirectPartReference]:
    """Load the indirect material reference parts with their tolerance bands."""
    global _CACHED_PARTS
    if _CACHED_PARTS is not None and not force_reload and not rebuild_from_excel:
        return _CACHED_PARTS

    if rebuild_from_excel or not JSON_CACHE_PATH.exists():
        if EXCEL_PATH.exists():
            raw_data = build_reference_from_excel(EXCEL_PATH)
        else:
            raise FileNotFoundError(f"Indirect reference file not found: {EXCEL_PATH}")
    else:
        with open(JSON_CACHE_PATH, "r", encoding="utf-8") as f:
            raw_data = json.load(f)

    parts: List[IndirectPartReference] = []
    for item in raw_data.get("parts", []):
        el_map: Dict[str, IndirectElementBand] = {}
        for sym, ed in item.get("elements", {}).items():
            ref_v = float(ed["reference_value"])
            tol, min_v, max_v = compute_tolerance_band(ref_v)
            el_map[sym] = IndirectElementBand(
                element=sym,
                reference_value=ref_v,
                tolerance=tol,
                min_value=min_v,
                max_value=max_v,
            )
        parts.append(
            IndirectPartReference(
                sn=int(item["sn"]),
                part_name=str(item["part_name"]),
                location=str(item.get("location", "")),
                material=str(item["material"]),
                elements=el_map,
            )
        )

    _CACHED_PARTS = parts
    return parts


DEFAULT_INDIRECT_FAMILY_LABELS: Dict[str, str] = {
    "SS 304": "Austenitic Stainless Steel (SS 304)",
    "SS 301": "Austenitic Stainless Steel (SS 301)",
    "Mn Steel": "High-Manganese Cr-Mn Austenitic Steel (Mn Steel)",
    "AiSi 410": "Martensitic / Cr Alloy Steel (AiSi 410)",
    "EN 31": "High-Carbon Chromium Bearing Steel (EN 31)",
    "CS": "Carbon Steel (CS)",
}


def _load_raw_indirect_json() -> Dict[str, Any]:
    if not JSON_CACHE_PATH.exists():
        if EXCEL_PATH.exists():
            return build_reference_from_excel(EXCEL_PATH)
        return {
            "version": "1.0.0",
            "source_file": "data/Cleaning area.xlsx",
            "scope": "Indirect Material Composition - Cleaning Area",
            "tolerance_rules": [
                {"range": "< 1", "tolerance": 0.25, "min_formula": "Value * 0.75", "max_formula": "Value * 1.25"},
                {"range": "1 to 5 (inclusive)", "tolerance": 0.20, "min_formula": "Value * 0.80", "max_formula": "Value * 1.20"},
                {"range": "> 5", "tolerance": 0.10, "min_formula": "Value * 0.90", "max_formula": "Value * 1.10"},
            ],
            "material_families": {},
            "parts": [],
        }
    return json.loads(JSON_CACHE_PATH.read_text(encoding="utf-8"))


def _rebuild_families_summary(raw_data: Dict[str, Any]) -> None:
    """Synchronize `material_families` dictionary with `parts` while preserving empty custom families."""
    existing_fams = raw_data.get("material_families", {})
    new_fams: Dict[str, List[str]] = {}
    for p in raw_data.get("parts", []):
        mat = str(p.get("material", "")).strip()
        pname = str(p.get("part_name", "")).strip()
        if mat:
            new_fams.setdefault(mat, [])
            if pname and pname not in new_fams[mat]:
                new_fams[mat].append(pname)
    # Preserve any empty user-created family
    for k in existing_fams:
        if k not in new_fams:
            new_fams[k] = []
    raw_data["material_families"] = new_fams


def get_indirect_kb_payload() -> Dict[str, Any]:
    """Return formatted Indirect Knowledge Base data for Frontend UI & REST API."""
    raw_data = _load_raw_indirect_json()
    custom_labels = raw_data.get("family_labels", {})
    labels_map = {**DEFAULT_INDIRECT_FAMILY_LABELS, **custom_labels}

    parts_formatted = []
    for p in raw_data.get("parts", []):
        el_list = []
        for sym, ed in p.get("elements", {}).items():
            ref_v = float(ed["reference_value"])
            tol, min_v, max_v = compute_tolerance_band(ref_v)
            el_list.append({
                "element": sym,
                "reference_value": round(ref_v, 4),
                "tolerance": tol,
                "tolerance_pct": int(round(tol * 100)),
                "tolerance_label": f"±{int(round(tol * 100))}%",
                "min_value": round(min_v, 4),
                "max_value": round(max_v, 4),
            })
        parts_formatted.append({
            "sn": int(p["sn"]),
            "part_name": p["part_name"],
            "location": p.get("location", ""),
            "material": p["material"],
            "material_label": labels_map.get(p["material"], p["material"]),
            "elements": p.get("elements", {}),
            "elements_list": el_list,
        })

    families_list = []
    mat_fams = raw_data.get("material_families", {})
    for fam_name, part_names in mat_fams.items():
        fam_parts = [pt for pt in parts_formatted if pt["material"] == fam_name]
        # Aggregate element min/max across parts in this family
        elem_ranges: Dict[str, Dict[str, float]] = {}
        for pt in fam_parts:
            for el_item in pt["elements_list"]:
                sym = el_item["element"]
                if sym not in elem_ranges:
                    elem_ranges[sym] = {
                        "min": el_item["min_value"],
                        "max": el_item["max_value"],
                        "ref_min": el_item["reference_value"],
                        "ref_max": el_item["reference_value"],
                        "count": 1,
                    }
                else:
                    elem_ranges[sym]["min"] = min(elem_ranges[sym]["min"], el_item["min_value"])
                    elem_ranges[sym]["max"] = max(elem_ranges[sym]["max"], el_item["max_value"])
                    elem_ranges[sym]["ref_min"] = min(elem_ranges[sym]["ref_min"], el_item["reference_value"])
                    elem_ranges[sym]["ref_max"] = max(elem_ranges[sym]["ref_max"], el_item["reference_value"])
                    elem_ranges[sym]["count"] += 1

        families_list.append({
            "name": fam_name,
            "label": labels_map.get(fam_name, fam_name),
            "parts_count": len(fam_parts),
            "part_names": part_names,
            "element_ranges": [
                {
                    "element": k,
                    "min_value": round(v["min"], 3),
                    "max_value": round(v["max"], 3),
                    "ref_min": round(v["ref_min"], 3),
                    "ref_max": round(v["ref_max"], 3),
                    "parts_support": v["count"],
                }
                for k, v in elem_ranges.items()
            ],
        })

    return {
        "version": raw_data.get("version", "1.0.0"),
        "source_file": raw_data.get("source_file", "data/Cleaning area.xlsx"),
        "scope": raw_data.get("scope", "Indirect Material Composition - Cleaning Area"),
        "tolerance_rules": raw_data.get("tolerance_rules", []),
        "family_labels": labels_map,
        "families": families_list,
        "parts": parts_formatted,
    }


def save_indirect_part_customization(payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Update an existing Indirect Part / Source (by `sn`) or create a new one,
    recalculating ±25% / ±20% / ±10% tolerance ranges for all populated elements.
    """
    global _CACHED_PARTS
    raw_data = _load_raw_indirect_json()
    parts = raw_data.setdefault("parts", [])

    part_name = str(payload.get("part_name") or payload.get("partName") or "").strip()
    if not part_name:
        raise ValueError("Indirect Part Name is required.")
    location = str(payload.get("location", "Cleaning Area")).strip()
    material = str(payload.get("material", "SS 304")).strip()
    if not material:
        raise ValueError("Material Family classification is required.")

    raw_elements = payload.get("elements", {})
    cleaned_elements: Dict[str, Dict[str, Any]] = {}
    if isinstance(raw_elements, dict):
        items_iter = raw_elements.items()
    elif isinstance(raw_elements, list):
        items_iter = [(x.get("element"), x.get("reference_value", x.get("referenceValue"))) for x in raw_elements if isinstance(x, dict)]
    else:
        items_iter = []

    for sym_raw, val_raw in items_iter:
        sym = str(sym_raw or "").strip()
        if not sym:
            continue
        if isinstance(val_raw, dict):
            val_raw = val_raw.get("reference_value", val_raw.get("referenceValue"))
        if val_raw is None or str(val_raw).strip() == "":
            continue
        try:
            ref_val = float(val_raw)
        except (ValueError, TypeError):
            continue
        if ref_val <= 0:
            continue
        tol, min_v, max_v = compute_tolerance_band(ref_val)
        cleaned_elements[sym] = {
            "element": sym,
            "reference_value": round(ref_val, 4),
            "tolerance": tol,
            "min_value": min_v,
            "max_value": max_v,
        }

    sn_raw = payload.get("sn")
    target_sn: Optional[int] = int(sn_raw) if sn_raw not in (None, "", 0, "0") else None

    updated = False
    if target_sn is not None:
        for p in parts:
            if int(p.get("sn", -1)) == target_sn:
                p["part_name"] = part_name
                p["location"] = location
                p["material"] = material
                p["elements"] = cleaned_elements
                updated = True
                break

    if not updated:
        next_sn = max([int(p.get("sn", 0)) for p in parts] + [0]) + 1
        target_sn = next_sn
        parts.append({
            "sn": target_sn,
            "part_name": part_name,
            "location": location,
            "material": material,
            "elements": cleaned_elements,
        })

    _rebuild_families_summary(raw_data)
    JSON_CACHE_PATH.write_text(json.dumps(raw_data, indent=2), encoding="utf-8")
    _CACHED_PARTS = None
    load_indirect_reference(force_reload=True)
    return get_indirect_kb_payload()


def delete_indirect_part_customization(sn: int) -> Dict[str, Any]:
    """Remove an indirect part by `sn` and refresh the Indirect Knowledge Base."""
    global _CACHED_PARTS
    raw_data = _load_raw_indirect_json()
    parts = raw_data.get("parts", [])
    raw_data["parts"] = [p for p in parts if int(p.get("sn", -1)) != int(sn)]
    _rebuild_families_summary(raw_data)
    JSON_CACHE_PATH.write_text(json.dumps(raw_data, indent=2), encoding="utf-8")
    _CACHED_PARTS = None
    load_indirect_reference(force_reload=True)
    return get_indirect_kb_payload()


def save_indirect_family_customization(
    old_name: Optional[str],
    new_name: str,
    label: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Rename an existing Indirect Material Family across all parts (or create a new one)
    and update its human-readable description label.
    """
    global _CACHED_PARTS
    new_clean = str(new_name or "").strip()
    if not new_clean:
        raise ValueError("Indirect Material Family name cannot be empty.")

    raw_data = _load_raw_indirect_json()
    family_labels = raw_data.setdefault("family_labels", dict(DEFAULT_INDIRECT_FAMILY_LABELS))
    family_aliases = raw_data.setdefault("family_aliases", {})
    mat_fams = raw_data.setdefault("material_families", {})

    old_clean = str(old_name or "").strip()
    if old_clean and old_clean != new_clean:
        # Update all parts referencing old_clean
        for p in raw_data.get("parts", []):
            if p.get("material") == old_clean:
                p["material"] = new_clean
        # Track canonical alias so Stage 1 rule classification maps seamlessly
        canonical_base = family_aliases.pop(old_clean, old_clean)
        family_aliases[new_clean] = canonical_base
        old_label = family_labels.pop(old_clean, old_clean)
        family_labels[new_clean] = str(label).strip() if label else old_label
        if old_clean in mat_fams:
            mat_fams[new_clean] = mat_fams.pop(old_clean)
    else:
        if new_clean not in mat_fams:
            mat_fams[new_clean] = []
        if label and str(label).strip():
            family_labels[new_clean] = str(label).strip()
        elif new_clean not in family_labels:
            family_labels[new_clean] = new_clean

    _rebuild_families_summary(raw_data)
    JSON_CACHE_PATH.write_text(json.dumps(raw_data, indent=2), encoding="utf-8")
    _CACHED_PARTS = None
    load_indirect_reference(force_reload=True)
    return get_indirect_kb_payload()

