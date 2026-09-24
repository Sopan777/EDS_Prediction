import os
from fastapi import UploadFile, File
from fastapi import FastAPI, Request, Depends, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from sqlalchemy import func, asc, desc
from datetime import datetime, date
import csv
import io
from openpyxl import load_workbook
from backend.database import engine, Base, get_db
from backend import models
from backend.auth import authenticate_user, create_user, get_current_user


# ================= CODE GENERATORS =================

def generate_record_code(date_str, db):
    date_obj = datetime.strptime(date_str, "%Y-%m-%d")
    date_part = date_obj.strftime("%Y%m%d")

    # Count records created on that date
    count = db.query(models.Record).filter(
        models.Record.date == date_obj.date()
    ).count()

    return f"NPR-{date_part}-{str(count + 1).zfill(4)}"


def generate_case_code(case_date, db):
    date_str = case_date.strftime("%Y%m%d")

    # Get highest existing number for this date
    last_case = db.query(models.Case).filter(
        models.Case.case_code.like(f"NPC-{date_str}-%")
    ).order_by(models.Case.case_code.desc()).first()

    if last_case and last_case.case_code:
        last_number = int(last_case.case_code.split("-")[-1])
        new_number = last_number + 1
    else:
        new_number = 1

    return f"NPC-{date_str}-{str(new_number).zfill(4)}"

"""
def generate_case_code(date_obj, db):
    date_part = date_obj.strftime("%Y%m%d")

    # Count cases created on that date
    count = db.query(models.Case).filter(
        models.Case.created_at == date_obj
    ).count()

    return f"NPC-{date_part}-{str(count + 1).zfill(4)}"
"""
app = FastAPI()

# Serve static files
app.mount("/static", StaticFiles(directory="static"), name="static")

Base.metadata.create_all(bind=engine)

templates = Jinja2Templates(directory="templates")

# ================= CURRENT EXCEL SOURCE =================

EXCEL_FILE = r"\\bosch.com\dfsrb\DfsIN\LOC\Na\DS\QMM\02_Projects\04_QMM3_Common\08_Projects_GAs\03_LAC\2024 - Snehal Yelmalle\CRIN Internal Rejection Analysis\CRIN Line rejection analysis updated.xlsx"


def load_current_excel():

    if not os.path.exists(EXCEL_FILE):

        return {
            "headers": [],
            "rows": [],
            "total_rows": 0
        }


    workbook = load_workbook(
        EXCEL_FILE,
        data_only=True,
        read_only=True
    )


    if "Particle Summary" in workbook.sheetnames:

        sheet = workbook["Particle Summary"]

    else:

        sheet = workbook[
            workbook.sheetnames[0]
        ]


    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )


    workbook.close()


    if not rows:

        return {
            "headers": [],
            "rows": [],
            "total_rows": 0
        }


    headers = []

    for value in rows[0]:

        if value is None:

            headers.append("")

        else:

            headers.append(
                str(value).strip()
            )


    data_rows = []


    for row in rows[1:]:

        cleaned_row = []


        for value in row:

            if value is None:

                cleaned_row.append("")


            elif hasattr(
                value,
                "strftime"
            ):

                cleaned_row.append(
                    value.strftime(
                        "%d-%m-%Y"
                    )
                )


            else:

                cleaned_row.append(
                    str(value).strip()
                )


        if any(
            str(value).strip()
            for value in cleaned_row
        ):

            data_rows.append(
                cleaned_row
            )


    return {
        "headers": headers,
        "rows": data_rows,
        "total_rows": len(data_rows)
    }


# ================= LOGIN =================

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse("login.html", {
        "request": request,
        "show_header": True,
        "hide_home_icon": True
    })


@app.post("/login")
def login(request: Request,
          username: str = Form(...),
          password: str = Form(...),
          db: Session = Depends(get_db)):

    user = authenticate_user(db, username, password)

    if not user:
        return templates.TemplateResponse("login.html", {
            "request": request,
            "error": "Invalid username or password.",
            "show_header": True,
            "hide_home_icon": True
        })

    response = RedirectResponse("/", status_code=302)
    response.set_cookie("user_id", str(user.id))
    return response


# ================= REGISTER =================

@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request):
    return templates.TemplateResponse("register.html", {
        "request": request,
        "show_header": True,
        "hide_home_icon": True
    })


@app.post("/register")
def register(request: Request,
             first_name: str = Form(...),
             last_name: str = Form(...),
             username: str = Form(...),
             email: str = Form(...),
             password: str = Form(...),
             confirm_password: str = Form(...),
             department: str = Form(...),
             role: str = Form(...),
             db: Session = Depends(get_db)):

    error = create_user(
        db,
        first_name,
        last_name,
        username,
        email,
        password,
        confirm_password,
        department,
        role
    )

    if error:
        return templates.TemplateResponse("register.html", {
            "request": request,
            "error": error,
            "show_header": True,
            "hide_home_icon": True,
            "form_data": {
                "first_name": first_name,
                "last_name": last_name,
                "username": username,
                "email": email,
                "department": department,
                "role": role
            }
        })

    return RedirectResponse("/login", status_code=302)

# ================= LOGOUT =================

@app.get("/logout")
def logout():
    response = RedirectResponse("/login")
    response.delete_cookie("user_id")
    return response


# ================= HOME =================

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    try:
        user = get_current_user(request)
    except:
        return RedirectResponse("/login")

    return templates.TemplateResponse("home.html", {
        "request": request,
        "user": user,
        "show_header": True
    })


# ================= DASHBOARD =================

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request,
              from_date: str = None,
              to_date: str = None,
              shift: str = None,
              line: str = None,
              auditor: str = None,
              type: str = None,
              db: Session = Depends(get_db)):

    try:
        user = get_current_user(request)
    except:
        return RedirectResponse("/login")

    query = db.query(models.Record)

    if user.role != "admin":
        query = query.filter(models.Record.owner_id == user.id)

    if from_date:
        query = query.filter(
            models.Record.date >=
            datetime.strptime(from_date, "%Y-%m-%d").date()
        )

    if to_date:
        query = query.filter(
            models.Record.date <=
            datetime.strptime(to_date, "%Y-%m-%d").date()
        )

    if shift:
        query = query.filter(
            models.Record.shift == shift
        )

    if line:
        query = query.filter(
            models.Record.line == line
        )

    if auditor:
        query = query.filter(
            models.Record.auditor_name == auditor
        )

    if type:
        query = query.filter(
            models.Record.type == type
        )

    records = query.order_by(
        models.Record.date.asc()
    ).all()


    # ================= KPI =================

    total_records = len(records)

    total_parts = sum(
        (r.parts_checked or 0)
        for r in records
    )

    total_nok = sum(
        (r.nok or 0)
        for r in records
    )

    nok_rate = round(
        (total_nok / total_parts * 100),
        2
    ) if total_parts else 0


    # ================= FILTER OPTIONS =================

    source_query = db.query(models.Record)

    if user.role != "admin":
        source_query = source_query.filter(
            models.Record.owner_id == user.id
        )

    source_records = source_query.all()

    shifts = sorted({
        r.shift
        for r in source_records
        if r.shift
    })

    lines = sorted({
        r.line
        for r in source_records
        if r.line
    })

    auditors = sorted({
        r.auditor_name
        for r in source_records
        if r.auditor_name
    })

    types = sorted({
        r.type
        for r in source_records
        if r.type
    })


    # ================= DAILY TREND =================

    daily = {}

    for r in records:

        key = (
            r.date.strftime("%d %b")
            if r.date
            else "Unknown"
        )

        if key not in daily:

            daily[key] = {
                "audits": 0,
                "nok": 0
            }

        daily[key]["audits"] += 1
        daily[key]["nok"] += r.nok or 0


    daily_labels = list(
        daily.keys()
    )

    daily_audits = [
        daily[k]["audits"]
        for k in daily_labels
    ]

    daily_nok = [
        daily[k]["nok"]
        for k in daily_labels
    ]


    # ================= SHIFT SUMMARY =================

    shift_summary = {}

    for r in records:

        if not r.shift:
            continue

        shift_summary.setdefault(
            r.shift,
            0
        )

        shift_summary[r.shift] += (
            r.nok or 0
        )


    shift_labels = list(
        shift_summary.keys()
    )

    shift_nok = [
        shift_summary[k]
        for k in shift_labels
    ]


    # ================= LINE SUMMARY =================

    line_summary = {}

    for r in records:

        if not r.line:
            continue

        line_summary.setdefault(
            r.line,
            0
        )

        line_summary[r.line] += (
            r.nok or 0
        )


    line_labels = list(
        line_summary.keys()
    )

    line_nok = [
        line_summary[k]
        for k in line_labels
    ]


    # ================= AUDITOR DAY-WISE =================

    auditor_daywise = {}

    for r in records:

        auditor_name = (
            r.auditor_name
            or "Unknown"
        )

        auditor_daywise.setdefault(
            auditor_name,
            {}
        )

        date_key = (
            r.date.strftime("%Y-%m-%d")
            if r.date
            else "Unknown"
        )

        auditor_daywise[
            auditor_name
        ][date_key] = (
            auditor_daywise[
                auditor_name
            ].get(date_key, 0) + 1
        )


    auditor_daywise = {

        auditor_name: [

            {
                "date": d,
                "count": count
            }

            for d, count
            in sorted(day_data.items())

        ]

        for auditor_name, day_data
        in sorted(
            auditor_daywise.items()
        )

    }


    # ================= RECENT NOK CASES =================

    case_query = (
        db.query(models.Case)
        .join(models.Record)
    )

    if user.role != "admin":

        case_query = case_query.filter(
            models.Record.owner_id == user.id
        )

    if from_date:

        case_query = case_query.filter(
            models.Record.date >=
            datetime.strptime(
                from_date,
                "%Y-%m-%d"
            ).date()
        )

    if to_date:

        case_query = case_query.filter(
            models.Record.date <=
            datetime.strptime(
                to_date,
                "%Y-%m-%d"
            ).date()
        )

    if shift:

        case_query = case_query.filter(
            models.Record.shift == shift
        )

    if line:

        case_query = case_query.filter(
            models.Record.line == line
        )

    if auditor:

        case_query = case_query.filter(
            models.Record.auditor_name == auditor
        )

    if type:

        case_query = case_query.filter(
            models.Record.type == type
        )


    recent_cases = (
        case_query
        .order_by(models.Case.id.desc())
        .limit(6)
        .all()
    )

    # Load the latest Excel data
    excel_data = load_current_excel()
    return templates.TemplateResponse(
        "dashboard.html",
        {

            "request": request,

            "user": user,

            "total_records":
                total_records,

            "total_parts":
                total_parts,

            "total_nok":
                total_nok,

            "nok_rate":
                nok_rate,

            "current_date":
                date.today().strftime(
                    "%d %b %Y"
                ),

            "auditor_daywise":
                auditor_daywise,

            "recent_cases":
                recent_cases,

            "daily_labels":
                daily_labels,

            "daily_audits":
                daily_audits,

            "daily_nok":
                daily_nok,

            "shift_labels":
                shift_labels,

            "shift_nok":
                shift_nok,

            "line_labels":
                line_labels,

            "line_nok":
                line_nok,

            "shifts":
                shifts,

            "lines":
                lines,

            "auditors":
                auditors,

            "types":
                types,

            "filters": {

                "from_date":
                    from_date,

                "to_date":
                    to_date,

                "shift":
                    shift,

                "line":
                    line,

                "auditor":
                    auditor,

                "type":
                    type
            },
"excel_headers":
    excel_data["headers"],

"excel_rows":
    excel_data["rows"],

"excel_total_rows":
    excel_data["total_rows"],

"show_header":
    True

        }
    )
    
# ================= ADD RECORD =================

@app.get("/add", response_class=HTMLResponse)
def add_page(request: Request):
    try:
        get_current_user(request)
    except:
        return RedirectResponse("/login")

    return templates.TemplateResponse("add_record.html", {
        "request": request,
        "show_header": True
    })


@app.post("/add")
def add_record(request: Request,
               date: str = Form(...),
               auditor_name: str = Form(...),
               shift: str = Form(...),
               line: str = Form(...),
               type: str = Form(...),
               parts_checked: int = Form(...),
               nok: int = Form(...),
               db: Session = Depends(get_db)):

    user = get_current_user(request)

    record = models.Record(
        record_code=generate_record_code(date, db),  # assuming this already exists
        date=datetime.strptime(date, "%Y-%m-%d"),
        auditor_name=auditor_name,
        shift=shift,
        line=line,
        type=type,
        parts_checked=parts_checked,
        nok=nok,
        owner_id=user.id
    )

    db.add(record)
    db.commit()
    db.refresh(record)

    # CREATE EMPTY CASE ROWS (NO case_code YET)
    if nok and nok > 0:
        for _ in range(nok):
            new_case = models.Case(
                record_id=record.id,
                created_at=record.date
            )
            db.add(new_case)

        db.commit()
        
        # Open first case only
        first_case = db.query(models.Case).filter(
            models.Case.record_id == record.id
        ).order_by(models.Case.id).first()

        return RedirectResponse(f"/edit-case/{first_case.id}", status_code=302)

        

    return RedirectResponse("/view", status_code=302)

# ================= EDIT RECORD =================

@app.get("/edit/{record_id}", response_class=HTMLResponse)
def edit_page(record_id: int, request: Request, db: Session = Depends(get_db)):

    user = get_current_user(request)

    record = db.query(models.Record).filter(models.Record.id == record_id).first()

    if not record:
        raise HTTPException(status_code=404)

    if user.role != "admin" and record.owner_id != user.id:
        raise HTTPException(status_code=403)

    return templates.TemplateResponse("add_record.html", {
        "request": request,
        "record": record,
        "show_header": True
    })


@app.post("/edit/{record_id}")
def edit_record(record_id: int,
                request: Request,
                date: str = Form(...),
                auditor_name: str = Form(...),
                shift: str = Form(...),
                line: str = Form(...),
                type: str = Form(...),
                parts_checked: int = Form(...),
                nok: int = Form(...),
                db: Session = Depends(get_db)):

    user = get_current_user(request)

    record = db.query(models.Record).filter(models.Record.id == record_id).first()

    if user.role != "admin" and record.owner_id != user.id:
        raise HTTPException(status_code=403)

    record.date = datetime.strptime(date, "%Y-%m-%d")
    record.auditor_name = auditor_name
    record.shift = shift
    record.line = line
    record.type = type
    record.parts_checked = parts_checked
    record.nok = nok

    db.commit()

    return RedirectResponse("/view", status_code=302)


# =======================================

# ================= EDIT CASES =================

@app.get("/edit-cases/{record_id}", response_class=HTMLResponse)
def edit_cases(record_id: int,
               request: Request,
               db: Session = Depends(get_db)):

    user = get_current_user(request)

    record = db.query(models.Record).filter(
        models.Record.id == record_id
    ).first()

    if not record:
        raise HTTPException(status_code=404)

    # Permission check
    if user.role != "admin" and record.owner_id != user.id:
        raise HTTPException(status_code=403)

    cases = db.query(models.Case).filter(
        models.Case.record_id == record_id
    ).all()

    return templates.TemplateResponse("edit_cases.html", {
        "request": request,
        "record": record,
        "cases": cases,
        "show_header": True
    })
    
@app.post("/update-case/{case_id}")
def update_case(case_id: int,
                request: Request,
                rejection_station: str = Form(...),
                particle_location: str = Form(...),
                body: str = Form(...),
                valve: str = Form(...),
                nozzle: str = Form(...),
                magnet: str = Form(...),
                remark: str = Form(""),
                bin_number: str = Form(""),
                image1: UploadFile = File(None),
                image2: UploadFile = File(None),
                image3: UploadFile = File(None),
                db: Session = Depends(get_db)):

    user = get_current_user(request)

    case = db.query(models.Case).filter(
        models.Case.id == case_id
    ).first()

    if not case:
        raise HTTPException(status_code=404)

    record = db.query(models.Record).filter(
        models.Record.id == case.record_id
    ).first()

    if user.role != "admin" and record.owner_id != user.id:
        raise HTTPException(status_code=403)

    # ✅ Generate case_code only once
    if not case.case_code:
        case.case_code = generate_case_code(case.created_at, db)

    # ---------------- Update fields ----------------
    case.rejection_station = rejection_station
    case.particle_location = particle_location
    case.body = body
    case.valve = valve
    case.nozzle = nozzle
    case.magnet = magnet
    case.remark = remark
    case.bin_number = bin_number

    # ---------------- Image Upload ----------------
    upload_folder = f"static/uploads/{case.case_code}"
    os.makedirs(upload_folder, exist_ok=True)

    def save_image(upload_file, filename):
        if upload_file and upload_file.filename:
            file_path = os.path.join(upload_folder, filename)
            with open(file_path, "wb") as buffer:
                buffer.write(upload_file.file.read())
            return file_path.replace("\\", "/")
        return None

    img1_path = save_image(image1, "image1.jpg")
    img2_path = save_image(image2, "image2.jpg")
    img3_path = save_image(image3, "image3.jpg")

    if img1_path:
        case.image1 = img1_path
    if img2_path:
        case.image2 = img2_path
    if img3_path:
        case.image3 = img3_path

    db.commit()

    # ================= SEQUENTIAL FLOW =================

    # Get all cases of this record ordered
    all_cases = db.query(models.Case).filter(
        models.Case.record_id == record.id
    ).order_by(models.Case.id).all()

    # Find current case index
    current_index = None
    for index, c in enumerate(all_cases):
        if c.id == case.id:
            current_index = index
            break

    # If next case exists → open it
    if current_index is not None and current_index + 1 < len(all_cases):
        next_case = all_cases[current_index + 1]
        return RedirectResponse(
            f"/edit-case/{next_case.id}",
            status_code=302
        )

    # If last case → go to Observations
    return RedirectResponse("/observations", status_code=302)
# ================= DELETE =================

@app.get("/delete/{record_id}")
def delete_record(record_id: int,
                  request: Request,
                  db: Session = Depends(get_db)):

    user = get_current_user(request)

    if user.role != "admin":
        raise HTTPException(status_code=403)

    record = db.query(models.Record).filter(models.Record.id == record_id).first()

    if record:
        db.delete(record)
        db.commit()

    return RedirectResponse("/view", status_code=302)


# ================= VIEW RECORDS =================

# ================= VIEW RECORDS =================

@app.get("/view", response_class=HTMLResponse)
def view_records(request: Request,
                 from_date: str = None,
                 to_date: str = None,
                 shift: str = None,
                 line: str = None,
                 sort_by: str = None,
                 order: str = "asc",
                 db: Session = Depends(get_db)):

    try:
        user = get_current_user(request)
    except:
        return RedirectResponse("/login")

    query = db.query(models.Record)

    if user.role != "admin":
        query = query.filter(models.Record.owner_id == user.id)

    if from_date:
        query = query.filter(models.Record.date >= datetime.strptime(from_date, "%Y-%m-%d"))

    if to_date:
        query = query.filter(models.Record.date <= datetime.strptime(to_date, "%Y-%m-%d"))

    if shift:
        query = query.filter(models.Record.shift == shift)

    if line:
        query = query.filter(models.Record.line == line)

    if sort_by in ["parts_checked", "nok"]:
        column = getattr(models.Record, sort_by)
        if order == "desc":
            query = query.order_by(desc(column))
        else:
            query = query.order_by(asc(column))

    records = query.all()

    return templates.TemplateResponse("view_records.html", {
        "request": request,
        "records": records,
        "user": user,
        "show_header": True,
        "filters": {
            "from_date": from_date,
            "to_date": to_date,
            "shift": shift,
            "line": line,
            "sort_by": sort_by,
            "order": order
        }
    })
# ================= DOWNLOAD =================

@app.get("/download", response_class=HTMLResponse)
def download_page(request: Request):
    try:
        get_current_user(request)
    except:
        return RedirectResponse("/login")

    return templates.TemplateResponse("download.html", {
        "request": request,
        "show_header": True
    })

@app.post("/download")
def download_csv(request: Request,
                 from_date: str = Form(None),
                 to_date: str = Form(None),
                 include_cases: str = Form(None),
                 db: Session = Depends(get_db)):

    user = get_current_user(request)

    query = db.query(models.Record)

    if user.role != "admin":
        query = query.filter(models.Record.owner_id == user.id)

    if from_date:
        query = query.filter(models.Record.date >= datetime.strptime(from_date, "%Y-%m-%d"))

    if to_date:
        query = query.filter(models.Record.date <= datetime.strptime(to_date, "%Y-%m-%d"))

    records = query.all()

    output = io.StringIO()
    writer = csv.writer(output)

    # ================= IF INCLUDE CASES =================
    if include_cases == "yes":

        writer.writerow([
            "Record Code", "Date", "Auditor", "Shift", "Line", "Type",
            "Parts Checked", "NOK",
            "Case Code", "Rejection Station", "Particle Location",
            "Body", "Valve", "Nozzle", "Magnet",
            "Remark", "Bin Number"
        ])

        for record in records:

            for case in record.cases:

                writer.writerow([
                    record.record_code,
                    record.date.strftime("%Y-%m-%d"),
                    record.auditor_name,
                    record.shift,
                    record.line,
                    record.type,
                    record.parts_checked,
                    record.nok,
                    case.case_code,
                    case.rejection_station,
                    case.particle_location,
                    case.body,
                    case.valve,
                    case.nozzle,
                    case.magnet,
                    case.remark,
                    case.bin_number
                ])

    # ================= RECORD ONLY =================
    else:

        writer.writerow([
            "Record Code", "Date", "Auditor", "Shift",
            "Line", "Type", "Parts Checked", "NOK"
        ])

        for record in records:
            writer.writerow([
                record.record_code,
                record.date.strftime("%Y-%m-%d"),
                record.auditor_name,
                record.shift,
                record.line,
                record.type,
                record.parts_checked,
                record.nok
            ])

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=records.csv"}
    )

# ================= ANALYTICS =================

@app.get("/analytics", response_class=HTMLResponse)
def analytics_page(request: Request,
                   analysis_type: str = "shift",
                   from_date: str = None,
                   to_date: str = None,
                   shift: str = None,
                   line: str = None,
                   db: Session = Depends(get_db)):

    try:
        user = get_current_user(request)
    except:
        return RedirectResponse("/login")


    allowed_types = {
        "shift",
        "line",
        "auditor_name",
        "type"
    }


    if analysis_type not in allowed_types:
        analysis_type = "shift"


    query = db.query(models.Record)


    if user.role != "admin":

        query = query.filter(
            models.Record.owner_id == user.id
        )


    if from_date:

        query = query.filter(
            models.Record.date >=
            datetime.strptime(
                from_date,
                "%Y-%m-%d"
            ).date()
        )


    if to_date:

        query = query.filter(
            models.Record.date <=
            datetime.strptime(
                to_date,
                "%Y-%m-%d"
            ).date()
        )


    if shift:

        query = query.filter(
            models.Record.shift == shift
        )


    if line:

        query = query.filter(
            models.Record.line == line
        )


    records = query.all()


    grouped = {}


    for r in records:

        key = (
            getattr(r, analysis_type)
            or "Unknown"
        )

        grouped.setdefault(
            key,
            {
                "parts": 0,
                "nok": 0
            }
        )


        grouped[key]["parts"] += (
            r.parts_checked or 0
        )


        grouped[key]["nok"] += (
            r.nok or 0
        )


    labels = list(
        grouped.keys()
    )


    parts = [
        grouped[k]["parts"]
        for k in labels
    ]


    nok = [
        grouped[k]["nok"]
        for k in labels
    ]


    return templates.TemplateResponse(
        "analytics.html",
        {

            "request":
                request,

            "labels":
                labels,

            "parts":
                parts,

            "nok":
                nok,

            "user":
                user,

            "show_header":
                True

        }
    )


@app.post("/analytics", response_class=HTMLResponse)
def analytics_data(request: Request,
                   analysis_type: str = Form(...),
                   from_date: str = Form(None),
                   to_date: str = Form(None),
                   shift: str = Form(None),
                   line: str = Form(None),
                   db: Session = Depends(get_db)):

    params = [
        f"analysis_type={analysis_type}"
    ]


    if from_date:
        params.append(
            f"from_date={from_date}"
        )


    if to_date:
        params.append(
            f"to_date={to_date}"
        )


    if shift:
        params.append(
            f"shift={shift}"
        )


    if line:
        params.append(
            f"line={line}"
        )


    return RedirectResponse(
        "/analytics?" + "&".join(params),
        status_code=303
    )
    
 
 # ================= OBSERVATIONS =================

@app.get("/observations", response_class=HTMLResponse)
def observations(request: Request,
                 from_date: str = None,
                 to_date: str = None,
                 shift: str = None,
                 line: str = None,
                 db: Session = Depends(get_db)):

    try:
        user = get_current_user(request)

    except:
        return RedirectResponse("/login")


    query = (
        db.query(models.Case)
        .join(models.Record)
    )


    if user.role != "admin":

        query = query.filter(
            models.Record.owner_id == user.id
        )


    if from_date:

        query = query.filter(
            models.Record.date >=
            datetime.strptime(
                from_date,
                "%Y-%m-%d"
            ).date()
        )


    if to_date:

        query = query.filter(
            models.Record.date <=
            datetime.strptime(
                to_date,
                "%Y-%m-%d"
            ).date()
        )


    if shift:

        query = query.filter(
            models.Record.shift == shift
        )


    if line:

        query = query.filter(
            models.Record.line == line
        )


    cases = (
        query
        .order_by(models.Case.id.desc())
        .all()
    )


    return templates.TemplateResponse(
        "observations.html",
        {

            "request":
                request,

            "cases":
                cases,

            "user":
                user,

            "show_header":
                True,

            "filters": {

                "from_date":
                    from_date,

                "to_date":
                    to_date,

                "shift":
                    shift,

                "line":
                    line

            }

        }
    )



@app.get("/delete-case/{case_id}")
def delete_case(case_id: int,
                request: Request,
                db: Session = Depends(get_db)):

    user = get_current_user(request)


    case = (
        db.query(models.Case)
        .filter(
            models.Case.id == case_id
        )
        .first()
    )


    if not case:

        raise HTTPException(
            status_code=404
        )


    record = (
        db.query(models.Record)
        .filter(
            models.Record.id ==
            case.record_id
        )
        .first()
    )


    if (
        user.role != "admin"
        and record.owner_id != user.id
    ):

        raise HTTPException(
            status_code=403
        )


    db.delete(case)

    db.commit()


    return RedirectResponse(
        "/observations",
        status_code=302
    )
@app.get("/debug/cases")
def debug_cases(db: Session = Depends(get_db)):
    cases = db.query(models.Case).all()
    return cases



@app.get("/debug/cases/{case_id}")
def debug_case(case_id: int, db: Session = Depends(get_db)):
    case = db.query(models.Case).filter(models.Case.id == case_id).first()
    return case



@app.get("/edit-case/{case_id}", response_class=HTMLResponse)
def edit_single_case(case_id: int,
                     request: Request,
                     db: Session = Depends(get_db)):

    user = get_current_user(request)

    case = db.query(models.Case).filter(
        models.Case.id == case_id
    ).first()

    if not case:
        raise HTTPException(status_code=404)

    record = db.query(models.Record).filter(
        models.Record.id == case.record_id
    ).first()

    if user.role != "admin" and record.owner_id != user.id:
        raise HTTPException(status_code=403)

    return templates.TemplateResponse("edit_case_single.html", {
        "request": request,
        "case": case,
        "record": record,
        "show_header": True
    })